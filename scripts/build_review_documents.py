"""Revise existing stage/technical reports without replacing their originals.

Uses bundled Python/python-docx. Experiment statistics are read, never invented.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import re
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
AUDIT=ROOT/'reports/final_audit_20261025'
EXP=AUDIT/'experiments'
OUT=ROOT/'deliverables/review_20261025'


def readcsv(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def number(x):return f'{float(x):.6f}'


def text_replace(p,text):
    # Some source figure paragraphs contain the picture and its caption.
    # Replacing Run.text would delete the drawing in that run.
    texts=p._p.xpath('.//w:t')
    anchor=texts[0] if texts else None
    while anchor is not None and anchor.getparent() is not p._p:anchor=anchor.getparent()
    formatting=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for run in p.runs:
        if run._r.xpath('.//w:drawing'):
            for node in run._r.xpath('.//w:t'):node.text=''
        else:run.text=''
    added=p.add_run(text)
    if formatting is not None:added._r.insert(0,formatting)
    if anchor is not None:anchor.addprevious(added._r)


def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for cell,txt in zip(t.rows[0].cells,headers):cell.text=str(txt)
    trpr=t.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row):cell.text=str(txt)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.size=Pt(9)
    return t


def addp(doc,text):
    p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(6)
    return p


def experiment_appendix(doc,metrics,tests,kind):
    doc.add_page_break()
    doc.add_heading('2026年10月25日算法验证与导师反馈补充',level=1)
    addp(doc,'本节为本次实际执行的补充验证。前文原留出成绩保留为历史实验，本节新成绩使用独立目录归档，不替换旧权重或旧图表。Web入口、页面、交互与后台集成源码保持不变。')
    if kind in ('final','w2model','w2weekly','dl','comparison','w3weekly','w3stage'):
        doc.add_heading('统一比较协议与时序回测',level=2)
        addp(doc,'RF、XGBoost和GRU使用完全相同的60,000条训练标签及目标时段；ML输入20维，GRU输入24×4连续历史与4维季节向量。统一训练样本量并不等于统一计算预算，GRU有内部验证早停，树模型沿用固定参数。三个外层评估窗为384–479、480–575、576–671；内部早停窗分别为336–383、432–479、528–575，训练目标从169开始至内部验证前一小时。全部2,880 beam参与每个外层评估，各276,480条。')
        addp(doc,'留出评估为672–839，共483,840条；训练目标169–587，内部验证588–671。每个模型的缩放器只用内部验证起点之前的raw观测拟合；已通过改变所有未来值而训练缩放参数不变的回归测试。固定seed42执行三折与留出，seed43仅重复留出；两种模型种子使用相同训练抽样（抽样seed42）。这是有限的初始化稳定性证据。')
        doc.add_paragraph('表A1 新实验留出集结果 seed42 原始KPI尺度')
        rows=[[r['KPI'],r['model'],number(r['MAE']),number(r['RMSE']),number(r['R2'])]
              for r in metrics if r['scenario']=='holdout' and r['seed']=='42']
        table(doc,['KPI','模型','MAE','RMSE','R²'],rows)
        addp(doc,'GRU在seed42的四项MAE均最低；DLThpVol仍由RF取得较低RMSE和较高R²。GRU适合作为MAE优先推荐，DLThpVol在关注大误差时保留RF。本结论针对本数据、配置和损失权重。历史传统综合排名1.417来自MAE/RMSE/R²名次等权，其中RMSE与R²同序，不能视为三个独立证据。')
        summary=readcsv(EXP/'walkforward_summary_flat.csv')
        doc.add_paragraph('表A2 新三折扩展窗口结果 seed42 均值及折间标准差')
        table(doc,['KPI','模型','MAE 均值±SD','RMSE 均值±SD'],
              [[r['KPI'],r['model'],f"{float(r['MAE_mean']):.4f} ± {float(r['MAE_std']):.4f}",
                f"{float(r['RMSE_mean']):.4f} ± {float(r['RMSE_std']):.4f}"] for r in summary])
        addp(doc,'三折均值中GRU的四项MAE/RMSE均低于本轮RF/XGB，但DLThpVol跨折RMSE波动较大，不能用一次留出结果概括所有时段。旧全量传统回测中DLThpVol的线性回归RMSE=0.967665低于RF=1.003675；该旧实验的训练预算不同，应独立解释。LSTM与Transformer本次没有新增三折或种子实验。')
        doc.add_heading('跨类别配对显著性与随机种子',level=2)
        addp(doc,'在相同beam和目标小时上，先计算逐记录绝对/平方误差，再按小时平均，形成168个配对损失。d为候选损失减参考损失，负值有利于候选。使用Newey–West HAC带宽24、双侧近似检验；每个场景和seed的3对模型×4KPI×2损失共24项一起Holm校正。95%区间为逐项HAC区间，未作同时区间校正。未为获得显著结果修改参数。')
        doc.add_paragraph('表A3 GRU对RF的留出配对损失检验 seed42')
        table(doc,['KPI及损失','平均差d','95% CI','Holm p'],
              [[r['KPI']+' '+r['loss'].replace('_loss',''),number(r['mean_loss_difference']),
                f"[{float(r['ci95_low']):.6f}, {float(r['ci95_high']):.6f}]",f"{float(r['holm_p']):.3g}"]
               for r in tests if r['seed']=='42' and r['reference']=='RandomForest' and r['contender']=='GRU'])
        addp(doc,'DLThpVol在seed42上的GRU对RF，MAE差−0.001853、MSE差+0.019228，两者Holm p均约0.354，不能判定差异显著。DLPRB、DLThpTime、MR_number的GRU对RF两类损失在该协议下均有显著改善。其余模型配对与seed43的完整效应、区间和p值见holdout_significance.csv，不将未显著解释为模型完全等效。')
        doc.add_paragraph('表A4 两组模型种子的留出误差 seed42及43')
        seedrows=[]
        for model in ('GRU','RandomForest','XGBoost'):
            for k in ('DLPRB','DLThpVol','DLThpTime','MR_number'):
                selected=[r for r in metrics if r['scenario']=='holdout' and r['model']==model and r['KPI']==k]
                selected.sort(key=lambda r:int(r['seed']))
                seedrows.append([k,model,' / '.join(f"{float(r['MAE']):.4f}" for r in selected),
                                 ' / '.join(f"{float(r['RMSE']):.4f}" for r in selected)])
        table(doc,['KPI','模型','MAE 42 / 43','RMSE 42 / 43'],seedrows)
        addp(doc,'seed43保留了GRU四项MAE较低、RF的DLThpVol RMSE较低的观察；两组种子不构成独立时间样本，不能把它们当作大样本稳定性证明。测试期已在历史选型中反复使用，本次属于固定协议复核，并非独立外部验证。')
    if kind in ('final','w2model','w2stats'):
        doc.add_heading('20维输入与超参数来源',level=2)
        table(doc,['KPI','五个输入字段'],[[k,' / '.join(f'{k}_z_{s}' for s in ('t','lag1','lag2','lag24','lag168'))]
                for k in ('DLPRB','DLThpVol','DLThpTime','MR_number')])
        addp(doc,'每组按当前t、滞后1、2、24、168小时排序。z为训练期StandardScaler无量纲值，监督目标为t+1原始KPI尺度。Time和base_station/cell/beam用于时间与身份对齐，不是额外输入维度。原数据未提供可据以换算的物理量纲，本报告保留数据集数值尺度，不擅写Mbps、百分比等单位。最长滞后要求169个连续整小时；训练与推理列顺序的合成测试已通过。')
        addp(doc,'主表传统ML及三种DL均使用固定人工设置；DL的Adam与早停属于权重优化及轮次选择，不等于自动超参数搜索。只有DLPRB的RF已有Optuna TPE 12次、256 beam、三折平均RMSE搜索；范围为树数40–160、深度8–24、最小叶样本5–50、max_features∈{sqrt,0.7,1.0}、max_samples 0.5–0.9。最佳137/13/9/sqrt/0.698071及CV RMSE=0.473700没有替换本次固定30树RF；未补造调优后的全量测试成绩。逐模型参数及定义/引用见参数登记簿。')
    if kind in ('final','anomaly','rootcause','w3weekly','w3stage','w2stats'):
        doc.add_heading('W4异常统计证据与联合规则',level=2)
        addp(doc,'复用既有W4统计，本次只重新核对全量标记逻辑和相关矩阵。483,840条唯一beam-hour记录中，3σ为22,059，IQR为51,318，IF为9,900，联合候选3,165。SVM仍为历史12,000条抽样测试，不参加共识投票。逐beam偏度/峰度支持偏态厚尾普遍存在，但阈值宽度、beam异质性和时间分布变化也影响标记率，不能把偏态指定为唯一因果机制。')
        shape=readcsv(ROOT/'reports/w4_anomaly_threshold_shape_summary.csv')
        table(doc,['KPI','|偏度|>1','超额峰度>3','3σ / IQR标记率'],
              [[r['KPI'],f"{float(r['fraction_abs_skewness_gt_1']):.2%}",f"{float(r['fraction_excess_kurtosis_gt_3']):.2%}",
                f"{float(r['3sigma_KPI_flag_rate']):.2%} / {float(r['IQR_KPI_flag_rate']):.2%}"] for r in shape])
        addp(doc,'零IQR回退分布数为0。逐KPI比例的分母各为483,840，整条beam-hour的总体标记则是任意KPI触发，不能把四个逐KPI比例相加当总体比例。')
        table(doc,['逻辑组合','候选数'],[['IF AND (3σ OR IQR)','3,165'],['IF AND IQR','3,159'],
              ['IF AND 3σ','2,080'],['IF AND 3σ AND IQR','2,074'],['IF OR 3σ OR IQR','58,108']])
        addp(doc,'OR允许任一统计阈值支持，AND要求IF与统计信号共同触发。本数据最终规则比IF AND IQR仅多6条，不能夸大OR分支的实际增益。候选规模的收紧/放宽不等于故障准确率或召回率改善；没有真实故障标签，不报告这些分类指标。')
        doc.add_heading('相关矩阵样本量与统计推断',level=2)
        corr=readcsv(ROOT/'reports/w4_anomaly_kpi_correlation_tests.csv')
        table(doc,['KPI组合','r','站级bootstrap 95% CI','Holm p'],
              [[r['KPI_1']+' / '+r['KPI_2'],f"{float(r['pearson_r']):.3f}",
                f"[{float(r['bootstrap_CI95_low']):.3f}, {float(r['bootstrap_CI95_high']):.3f}]",
                f"{float(r['holm_adjusted_p_value']):.3g}"] for r in corr])
        addp(doc,'每对n=3,165是候选记录数；30是基站聚类数，不代表3,165条独立样本。p值来自站级聚类稳健零回归斜率检验，六对一起Holm校正；区间另用站级bootstrap估计。六对校正后p均<0.05，仅说明已筛选候选中的关联。候选选择可能改变相关结构，单向站级聚类未完全处理跨站同小时共同冲击；30聚类仍为近似推断，不构成故障根因或因果证明。')
    doc.add_heading('代码修复与最新验收',level=2)
    addp(doc,'修复仅涉及底层小时数据契约、DL输出预检和scaler测试提取。训练加载器与推理拒绝共同缺小时、重复或非整数小时；保留ML 20维、DL窗口、逆标准化和IF AND (3σ OR IQR)定义。DL当前代码已仅在内部训练前缀拟合scaler；本次锁定该行为，没有把历史成绩改写为新成绩。')
    addp(doc,'本次既有自动化基线48项通过；新增算法及统计回归后最新完整测试57项通过（final_tests.txt）。六个Streamlit页面AppTest加载及切换均无异常，实际新模型回载及CSV导出通过，原异常模型5,760条检测通过。另用1轮、512条样本执行现有DL CLI仅作功能验收，不把其分数纳入模型比较。Web页面及服务源码哈希保持不变。本次未以真人浏览器重走每个上传/下载控件或实时WebSocket训练全流程，不能将自动验收冒充完整手工验收。')
    addp(doc,'证据根目录：reports/final_audit_20261025。新实验配置protocol.json、逐次runs.json、60组metrics.csv、hourly_losses.csv、48组holdout_significance.csv；实际权重及逐beam-hour配对预测位于resources/final_audit_20261025（Git忽略，需本机备份）。历史reports/和figures/数值保持原样。没有独立外部数据、全模型调参或生产实时网管接入，本次不扩大该范围。')


def main(source_root: Path):
    OUT.mkdir(parents=True,exist_ok=True)
    metrics=readcsv(EXP/'metrics.csv');tests=readcsv(EXP/'holdout_significance.csv')
    historic_comparison=readcsv(ROOT/'reports/week3/w3_ml_dl_comparison_summary.csv')
    # Flatten multi-index CSV for Word tables without depending on pandas runtime.
    with (EXP/'walkforward_summary.csv').open(encoding='utf-8',newline='') as f:
        rows=list(csv.reader(f))
    flat=[]
    for row in rows[3:]:
        flat.append(dict(model=row[0],KPI=row[1],MAE_mean=row[2],MAE_std=row[3],RMSE_mean=row[4],RMSE_std=row[5],R2_mean=row[6],R2_std=row[7]))
    with (EXP/'walkforward_summary_flat.csv').open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(flat[0]));w.writeheader();w.writerows(flat)
    sources=[
      ('W2_数据集统计分析报告',source_root / 'week2/w2-1.docx','w2stats'),
      ('W2_模型性能评估报告',source_root / 'week2/w2-2.docx','w2model'),
      ('W2_周报',source_root / 'week2/W2周报.docx','w2weekly'),
      ('W3_深度学习模型性能评估报告',source_root / 'week3/W3报告/01_深度学习模型性能评估报告_修订版.docx','dl'),
      ('W3_传统ML与DL模型对比分析',source_root / 'week3/W3报告/02_传统ML与DL模型对比分析_修订版.docx','comparison'),
      ('W3_异常检测结果报告',source_root / 'week3/W3报告/03_异常检测结果报告_修订版.docx','anomaly'),
      ('W3_异常根因分析初步方案',source_root / 'week3/W3报告/04_异常根因分析初步方案_修订版.docx','rootcause'),
      ('W3_第三阶段周实验报告',source_root / 'week3/W3报告/05_第三阶段周实验报告.docx','w3stage'),
      ('W3_周报',source_root / 'week3/W3报告/w3周报.docx','w3weekly'),
      ('完整项目技术报告',source_root / 'week4/5G网络KPI预测与异常检测系统_完整项目技术报告_可编辑.docx','final')]
    manifest=[]
    for name,source,kind in sources:
        source=Path(source);doc=Document(source)
        # Date and status corrections in place; historical numbers remain intact.
        for p in doc.paragraphs:
            old=p.text;new=old
            if any(phrase in old for phrase in ['尚未进行深度学习模型的滚动','未对深度学习模型执行同样','目前并无“GRU','尚未建立滚动回测','缺少独立外部测试、滚动回测']):
                new='历史阶段状态：'+old+' 2026年10月25日已增加GRU三折及RF/XGB/GRU跨类配对检验，详见本报告末尾补充验证；独立外部测试仍未完成。'
            if '后续正式滚动验证应按训练子段单独拟合scaler' in new.replace(' ',''):
                new+=' 当前代码已在内部训练前缀重拟合scaler，本次新实验独立保存，旧成绩仍属历史记录。'
            if kind=='final' and '先完成新版回归' in new:
                new=new.replace('先完成新版回归','按本次自动回归记录交付')
            if kind in ('w3weekly','w3stage') and ('下一周工作' in new or '下一阶段工作' in new):
                new=new.replace('下一周工作','历史阶段的后续安排').replace('下一阶段工作','历史阶段的后续安排')
            if kind in ('w3weekly','w3stage') and '开发交互界面' in new:
                new='历史阶段计划（本次已核验六模块集成及自动回归）：'+new
            if new!=old:text_replace(p,new)
        # The stage summary's original comparison table contained MAE only.
        # Add paired RMSE columns directly, rather than relying on an appendix.
        for t in doc.tables:
            header=' '.join(c.text for c in t.rows[0].cells)
            if 'MAE' in header and 'RMSE' not in header and ('相对' in header or '改善' in header):
                t.add_column(Inches(.9));t.add_column(Inches(.9))
                t.rows[0].cells[-2].text='传统最优RMSE'
                t.rows[0].cells[-1].text='DL最优RMSE及模型'
                for row in t.rows[1:]:
                    k=row.cells[0].text.strip()
                    found=next((r for r in historic_comparison if r['KPI']==k and r['metric']=='RMSE'),None)
                    if found:
                        row.cells[-2].text=f"{float(found['ML_best_score']):.4f}\n{found['ML_best_model']}"
                        row.cells[-1].text=f"{float(found['DL_best_score']):.4f}\n{found['DL_best_model']}"
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        old=p.text
                        new=old.replace('具备交互证据；全回归待记录','本次57项自动测试及六页AppTest通过；完整手工回归未重做')
                        new=new.replace('功能闭环已验收；新版完整回归和资源基准待单独复核','功能已集成；本次自动回归通过；完整手工回归及资源专项未重做')
                        if old!=new:text_replace(p,new)
        p=doc.paragraphs[0].insert_paragraph_before('2026年10月25日核验修订版  前文历史成绩保留 新实验及最新状态见末节补充验证')
        p.paragraph_format.space_after=Pt(8)
        # Keep existing artwork and structure; normalize only final report skipped numbering.
        if kind=='final':
            for prefix in ('图','表'):
                oldnums=[]
                for p in doc.paragraphs:
                    m=re.match(r'^'+prefix+r'\s*(\d+(?:-\d+)?)\s',p.text)
                    if m:oldnums.append(m.group(1))
                mapping={n:str(i+1) for i,n in enumerate(dict.fromkeys(oldnums))}
                for p in doc.paragraphs:
                    text=p.text
                    new=re.sub(prefix+r'\s*(\d+(?:-\d+)?)(?![\d-])',lambda m: prefix+' '+mapping.get(m.group(1),m.group(1)),text)
                    if new!=text:text_replace(p,new)
        # Add adjacent, specific figure context without replacing existing conclusions.
        for p in list(doc.paragraphs):
            if re.match(r'^图\s*\d+\s+',p.text):
                context='图注补充：历史gotsf-ds留出目标Time 672–839，n=483,840 beam-hour；误差使用原KPI尺度，均值曲线不替代逐样本评分。'
                if any(k in p.text for k in ['相关','矩阵']):context='图注补充：Pearson r，正式联合候选n=3,165、30基站聚类；六对Holm检验与区间见本次统计补充，相关不代表因果。'
                elif 'R²' in p.text or 'R2' in p.text:context='图注补充：历史留出Time 672–839，n=483,840逐beam-hour；R²为无量纲方差解释指标，越高越好，不是原KPI单位的误差。'
                elif '相对' in p.text:context='图注补充：同一483,840条历史留出记录，纵轴为相对误差降低百分比；正值改善，负值退步，分母为相应KPI及指标的传统ML最佳误差。'
                elif any(k in p.text for k in ['小时','时间索引']):context='图注补充：483,840条历史测试记录中筛出3,165条共识候选，按目标小时计数；每小时分母2,880 beam，最高Time 692有71条。'
                elif any(k in p.text for k in ['异常','标记']):context='图注补充：历史全量分母483,840 beam-hour；SVM另为12,000抽样；候选率不代表真实故障率。'
                elif any(k in p.text for k in ['验证','收敛','MSE变化']):context='图注补充：历史内部验证抽样12,000条，纵轴为标准化MSE、横轴为epoch；早停反映该次验证表现，新协议单独列示。'
                elif any(k in p.text for k in ['Frame','Web','界面','GTP','零值']):continue
                if 'RMSE' in p.text and '比较' in p.text:
                    context+=' 图示结论：历史留出集RF在四项KPI的RMSE均最低；新协议成绩另见附表。'
                elif 'MAE' in p.text and ('比较' in p.text or '对比' in p.text) and kind in ('w2weekly','w2model'):
                    context+=' 图示结论：历史留出集XGB在DLPRB、DLThpTime、MR_number的MAE较低，DLThpVol以RF最低。'
                text_replace(p,p.text+'\n'+context)
        experiment_appendix(doc,metrics,tests,kind)
        # Repair source tables extending beyond printable margins, retaining
        # column proportions. Do not delete illustrations or shrink body text.
        available=int(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)
        for t in doc.tables:
            widths=[int(c.width or 1) for c in t.columns]
            t.autofit=False
            t.alignment=WD_TABLE_ALIGNMENT.CENTER
            fitted=[int(w/sum(widths)*available) for w in widths]
            for col,w in zip(t.columns,fitted):col.width=w
            for row in t.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if len(row.cells)==len(fitted):
                    for cell,w in zip(row.cells,fitted):cell.width=w
            width=t._tbl.tblPr.find(qn('w:tblW'))
            if width is not None:
                width.set(qn('w:type'),'dxa');width.set(qn('w:w'),str(available//635))
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in t.rows[0].cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
            ind=t._tbl.tblPr.find(qn('w:tblInd'))
            if ind is not None:ind.set(qn('w:w'),'0')
        if kind=='w3weekly':
            # Source includes redundant body manual breaks that produced a
            # completely blank page after the expanded figure captions.
            original_paragraphs=doc.paragraphs[:43]
            first_break=True
            for p in original_paragraphs:
                breaks=p._p.xpath('.//w:br[@w:type="page"]')
                if not breaks:continue
                if first_break:first_break=False;continue
                for node in breaks:node.getparent().remove(node)
        assert len(doc.inline_shapes)==len(Document(source).inline_shapes), 'Source figure lost'
        out=OUT/(name+'_核验修订20261025.docx');doc.save(out)
        manifest.append(dict(source=str(source),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                             revised=str(out.relative_to(ROOT)),kind=kind,sha256=hashlib.sha256(out.read_bytes()).hexdigest()))
    (AUDIT/'report_revision_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Revised {len(manifest)} existing reports, originals unchanged')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path,
                        default=ROOT / 'docs/source_reports')
    args = parser.parse_args()
    main(args.source_root)
